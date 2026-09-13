"""Create synthetic Office fixtures with no user or business data."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

root = Path(__file__).parent
word = Document()
word.add_heading("ClawMaster Office 验收", 0)
word.add_paragraph("WORD_ORIGINAL_20260913")
table = word.add_table(rows=2, cols=2)
for cell, text in zip([c for row in table.rows for c in row.cells], ["项目", "数量", "测试", "8"]):
    cell.text = text
word.save(root / "document.docx")
excel = Workbook()
sheet = excel.active
sheet["A1"] = "EXCEL_ORIGINAL_20260913"
sheet["C2"] = 8
sheet["D2"] = 12
sheet["E2"] = "=C2*D2"
excel.save(root / "workbook.xlsx")
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[0])
slide.shapes.title.text = "ClawMaster Office 验收"
slide.placeholders[1].text = "POWERPOINT_ORIGINAL_20260913"
presentation.save(root / "presentation.pptx")
